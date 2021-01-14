from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart

wb = load_workbook("sample.xlsx")
ws = wb.active

# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart()  # 차트 종류 설정 (bar, Line, Pie ...)
# bar_chart.add_data(bar_value)  # 차트 데이텇 추가

# ws.add_chart(bar_chart, "E1")  # 차트 넣을 위치 정의
# wb.save("sample_chart.xlsx")

line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True)  # 차트 계열 설정
line_chart.title = " 성적표"  # 제목
line_chart.style = 10  # 미리 정의된 스타일을 적용, 사용자가 개별 지정 가능
line_chart.y_axis.title = "점수"
line_chart.x_axis.title = "번호"
ws.add_chart(bar_chart, "E1")

wb.save("sample_line_chart.xlsx")
